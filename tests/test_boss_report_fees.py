"""老板报表插入手续费列后的金额边界及跨页引用回归。"""

import unittest
import os
import tempfile
import zipfile

import openpyxl

from crawlers.generate_boss_report import (
    _TEMPLATE_PATH,
    _formula_text,
    _remap_summary_formula,
    _remove_empty_region_sections,
    _remove_empty_hong_kong_sections,
    _rewrite_summary_roster,
    add_summary_fee_layout,
    normalize_income_data_bars,
)


class BossReportFeeTests(unittest.TestCase):
    def setUp(self):
        self.wb = openpyxl.load_workbook(_TEMPLATE_PATH)
        self.total = _rewrite_summary_roster(
            self.wb['门店汇总'], ['香港测试店', '内地测试店'], {},
        )
        add_summary_fee_layout(self.wb, self.total, 2, 51)
        self.ws = self.wb['门店汇总']

    def tearDown(self):
        self.wb.close()

    def test_income_excludes_fees_and_preserves_kpay_adjustment(self):
        self.assertEqual(
            self.ws['Z2'].value,
            '=SUM(D2,F2,H2,J2,L2,N2,P2,R2,T2,V2,W2,X2,AB2,AC2,AD2,AE2)',
        )
        self.assertTrue(_formula_text(self.ws['P2'].value).endswith('-X2'))
        self.assertIn("'源数据'!$Q$2:$Q$50", self.ws['E2'].value)
        self.assertIn("'源数据'!$CA$2:$CA$50", self.ws['S2'].value)
        self.assertIn("'源数据'!$Y$2:$Y$50", self.ws['Y2'].value)
        self.assertEqual(self.ws['Q2'].value, '=IF(P2=0,0,"未提供")')
        data_bars = [
            (str(cf.sqref), rule.dataBar)
            for cf in self.ws.conditional_formatting
            for rule in self.ws.conditional_formatting._cf_rules[cf]
            if rule.type == 'dataBar'
        ]
        self.assertEqual({item[0] for item in data_bars}, {'D2:W3', 'X2:X3'})
        for _, data_bar in data_bars:
            self.assertEqual(data_bar.cfvo[0].type, 'num')
            self.assertEqual(data_bar.cfvo[0].val, 0.0)
            self.assertEqual(data_bar.cfvo[1].val, 'INDEX($Z:$Z,ROW())')

    def test_summary_keeps_two_blank_rows_and_subtracts_fees_once(self):
        for row in (self.total + 1, self.total + 2):
            self.assertTrue(all(c.value is None for c in self.ws[row]))
        fee, net = self.total + 3, self.total + 4
        self.assertEqual(self.ws[f'A{fee}'].value, '手续费汇总')
        self.assertEqual(self.ws[f'A{net}'].value, '实际收入')
        self.assertEqual(self.ws[f'Z{fee}'].value, f'=SUM(E{self.total},G{self.total},I{self.total},K{self.total},M{self.total},O{self.total},Q{self.total},S{self.total},U{self.total},Y{self.total})')
        self.assertEqual(self.ws[f'Z{net}'].value, f'=Z{self.total}-Z{fee}')
        self.assertIn('0.00', self.ws[f'Z{net}'].number_format)

    def test_each_store_has_net_income_immediately_after_gross_income(self):
        self.assertEqual(self.ws['Z1'].value, '收入汇总')
        self.assertEqual(self.ws['AA1'].value, '实际收入')
        self.assertEqual(self.ws['AA2'].value, '=Z2-SUM(E2,G2,I2,K2,M2,O2,Q2,S2,U2,Y2)')
        self.assertEqual(self.ws['AA3'].value, '=Z3-SUM(E3,G3,I3,K3,M3,O3,Q3,S3,U3,Y3)')
        # 合计只汇总实际数据行（2..门店数+1），与其他列一致，
        # 不包含数据行与合计行之间的空白备用行。
        self.assertEqual(self.ws[f'AA{self.total}'].value, '=SUM(AA2:AA3)')
        self.assertIn('0.00', self.ws['AA2'].number_format)
        self.assertIsNone(self.ws['AA4'].value)

    def test_ranking_overview_regions_and_source_references_stay_aligned(self):
        ranking = self.wb['门店排名']
        self.assertIn('门店汇总!Z:Z', _formula_text(ranking['E2'].value))
        self.assertIn('门店汇总!AF:AF', _formula_text(ranking['L2'].value))
        self.assertEqual(self.wb['老板总览']['A6'].value, "=SUM('门店汇总'!$Z$2:$Z$28)")
        # 货款比% = 总货款(AH) / 收入汇总(Z)
        self.assertEqual(self.ws['AI2'].value, '=IFERROR(AH2/Z2,0)')
        self.assertEqual(self.wb['内地门店']['P1'].value, '收入汇总')
        # 门店汇总 D 列（乐摇摇）= 源数据 非现金(P) + 现金(AO)
        self.assertIn('源数据!P:P', _formula_text(self.ws['D2'].value))
        self.assertIn('源数据!AO:AO', _formula_text(self.ws['D2'].value))
        self.assertEqual(self.ws['Z1'].value, '收入汇总')
        # 模板 Q..T 插入油菜花后，积分货款由 Q(17) 右移至 AF(32)
        self.assertEqual(self.ws['AF1'].value, '积分货款')
        self.assertEqual(self.ws['AB1'].value, '油菜花现金')
        self.assertEqual(self.ws['AE1'].value, '油菜花盈客宝')
        self.assertEqual(self.ws['E1'].value, '乐摇摇\n手续费')
        self.assertIn(f'A{self.total}:C{self.total}', str(self.ws.merged_cells))

    def test_formula_remap_preserves_literals_absolute_refs_and_other_sheets(self):
        formula = '=IF(\'门店汇总\'!$P2=0,"P2",SUM(源数据!P:P,P2,$S$2))'
        self.assertEqual(
            _remap_summary_formula(formula, '门店汇总'),
            '=IF(\'门店汇总\'!$Z2=0,"P2",SUM(源数据!P:P,Z2,$AD$2))',
        )

    def test_income_data_bars_are_solid_and_zero_based(self):
        handle, path = tempfile.mkstemp(prefix='workbuddy-databar-', suffix='.xlsx')
        os.close(handle)
        try:
            self.wb.save(path)
            self.wb.close()
            normalize_income_data_bars(path)
            with zipfile.ZipFile(path) as archive:
                xml = b''.join(
                    archive.read(name)
                    for name in archive.namelist()
                    if name.startswith('xl/worksheets/') and name.endswith('.xml')
                ).decode('utf-8')
            self.assertGreater(xml.count('<dataBar '), 0)
            self.assertNotIn('<dataBar>', xml)
            self.assertEqual(xml.count('gradient="0"'), xml.count('<dataBar '))
            self.assertEqual(xml.count('axisPosition="none"'), xml.count('<dataBar '))
        finally:
            if os.path.exists(path):
                os.unlink(path)

    def test_empty_hong_kong_sections_are_removed_without_broken_overview_refs(self):
        _remove_empty_hong_kong_sections(self.wb)

        self.assertNotIn('香港门店', self.wb.sheetnames)
        self.assertNotIn('香港排名', self.wb.sheetnames)
        overview = self.wb['老板总览']
        self.assertTrue(overview.row_dimensions[13].hidden)
        self.assertTrue(all(overview.cell(13, column).value is None for column in range(1, 11)))
        remaining_formulas = [
            cell.value
            for worksheet in self.wb.worksheets
            for row in worksheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith('=')
        ]
        self.assertFalse(any('香港门店' in formula for formula in remaining_formulas))
        self.assertFalse(any('香港排名' in formula for formula in remaining_formulas))

    def test_empty_region_sections_remove_redundant_mainland_pages(self):
        _remove_empty_region_sections(
            self.wb,
            ('内地门店', '内地排名', '香港门店', '香港排名 '),
            ('内地', '香港'),
        )

        for sheet_name in ('内地门店', '内地排名', '香港门店', '香港排名 '):
            self.assertNotIn(sheet_name, self.wb.sheetnames)
        overview = self.wb['老板总览']
        self.assertTrue(overview.row_dimensions[12].hidden)
        self.assertTrue(overview.row_dimensions[13].hidden)
        self.assertTrue(all(overview.cell(12, column).value is None for column in range(1, 11)))
        self.assertTrue(all(overview.cell(13, column).value is None for column in range(1, 11)))


if __name__ == '__main__':
    unittest.main()
