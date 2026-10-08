"""历史收入、目标月份与未知数据的回归验证。"""
import json
from contextlib import ExitStack
from unittest import mock

import pytest
from core import dashboard_data as dashboard, db, targets


def calculate(snapshots, active=("A", "B"), store_targets=None, scope=None, issues=()):
    with ExitStack() as stack:
        def load(day, venue_scope=None):
            value = snapshots.get(day)
            if value is None:
                return None
            return {k: v for k, v in value.items() if venue_scope is None or k in venue_scope}
        stack.enter_context(mock.patch.object(dashboard, "_load_venue_income", side_effect=load))
        stack.enter_context(mock.patch.object(dashboard, "_active_venue_set", return_value=set(active)))
        stack.enter_context(mock.patch.object(dashboard, "load_store_targets", return_value=store_targets or {}))
        stack.enter_context(mock.patch.object(dashboard, "load_store_regions", return_value={"A": "一区", "B": "一区", "closed": "二区"}))
        stack.enter_context(mock.patch.object(dashboard, "_snapshot_issues", return_value=list(issues)))
        stack.enter_context(mock.patch.object(dashboard, "enrich_dashboard", side_effect=lambda result, *args: result))
        return dashboard.get_dashboard_data("2026-09-07", scope)


SNAPSHOTS = {
    "2026-09-05": {"A": 60, "B": 20, "closed": 10},
    "2026-09-06": {"A": 80, "B": 30, "closed": 10},
    "2026-09-07": {"A": 100, "B": 40, "closed": 10},
    "2026-08-07": {"A": 200, "B": 100, "closed": 700},
}


def test_closed_history_is_in_totals_but_not_operating_rankings():
    result = calculate(SNAPSHOTS, store_targets={"A": 3000})
    assert result["curr_month_total"] == 150
    assert result["last_month_total"] == 1000
    assert result["month_change_pct"] == -85
    assert result["today_total"] == 30
    assert result["avg_daily"] == pytest.approx(150 / 7)
    assert result["daily_trend"][-1]["cumulative"] == 150
    assert result["daily_trend"][-1]["daily"] == 30
    assert {v for v, _ in result["top_cn"]} == {"A", "B"}
    assert result["regions"] == {"一区": 140, "二区": 10}
    assert result["completion_rate"] == 3.33
    assert result["target_actual"] == 100
    assert result["target_daily"] == pytest.approx(0.01)
    assert result["days_in_month"] == 30
    assert result["missing_target_venues"] == ["B"]
    assert result["region_completion"][0]["actual"] == 100


def test_missing_month_target_is_unknown_not_fallback_or_zero():
    result = calculate(SNAPSHOTS)
    for field in ["month_target", "completion_rate", "target_daily"]:
        assert result[field] is None
    assert result["region_completion"] == []
    assert result["target_month"] == "2026-09"
    assert not result["store_targets_loaded"]


def test_permission_scope_applies_to_history_and_trend():
    result = calculate(SNAPSHOTS, scope=(v for v in ["A"]), store_targets={"A": 3000, "B": 9000})
    assert result["curr_month_total"] == 100
    assert result["last_month_total"] == 200
    assert result["daily_trend"][-1]["cumulative"] == 100
    assert result["month_target"] == 0.3
    assert result["regions"] == {"一区": 100}


def test_missing_comparison_is_not_zero_change():
    result = calculate({k: v for k, v in SNAPSHOTS.items() if k not in {"2026-08-07", "2026-09-05"}})
    for field in ["last_month_total", "month_change_pct", "month_change", "yesterday_total", "day_change_pct"]:
        assert result[field] is None
    assert result["curr_month_total"] == 150


def test_zero_comparison_has_no_percentage():
    snapshots = dict(SNAPSHOTS, **{"2026-08-07": {"A": 0}, "2026-09-05": SNAPSHOTS["2026-09-06"]})
    result = calculate(snapshots)
    assert result["month_change"] == 150
    assert result["month_change_pct"] is None
    assert result["day_change_pct"] is None


def test_wrong_period_cannot_be_shown_as_monthly_income():
    result = calculate(SNAPSHOTS, issues=[{"date": "2026-09-06", "blocking": True, "reason": "range_mismatch"}])
    assert result["curr_month_total"] == 150
    assert result["today_total"] is None
    assert result["day_change"] is None
    assert result["daily_trend"][-1]["daily"] is None
    assert result["daily_trend"][-2]["cumulative"] is None


def test_target_file_matches_requested_month_even_if_configuration_is_old(tmp_path, monkeypatch):
    old = tmp_path / "2026-08_目标.xlsx"
    new = tmp_path / "2026-09_目标.xlsx"
    old.touch()
    monkeypatch.setattr(targets, "_TARGETS_DIR", str(tmp_path))
    monkeypatch.setattr(targets, "_config_path", lambda *args: str(old))
    assert targets.find_target_file("2026-09") == ""
    assert targets.load_store_targets("2026-09") == {}
    new.touch()
    assert targets.find_target_file("2026-09") == str(new)
    assert targets.find_target_file("2026-08") == str(old)
    assert targets.find_target_file() == str(old)
    assert targets.find_target_file("2027-09") == ""


def test_quality_checks_pairs_periods_and_permissions(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", str(tmp_path / "app.db"))
    db.init_db()
    conn = db.get_connection()
    rows = [("2026-09-01", "A", "p", "2026-09-01"),
            ("2026-09-01", "A", "q", "2026-09-01"),
            ("2026-09-02", "A", "p", "2026-09-02"),
            ("2026-09-02", "private", "q", None)]
    conn.executemany("INSERT INTO daily_summary(id,date,venue,platform,metrics_json,period_start) VALUES(?,?,?,?,?,?)",
                     [(str(i), day, venue, platform, "{}", start) for i, (day, venue, platform, start) in enumerate(rows)])
    conn.commit()
    conn.close()
    issues = dashboard._snapshot_issues(["2026-09-02"], {"A"})
    assert {(i["platform"], i["reason"]) for i in issues} == {("q", "venue_platform_missing"), ("p", "range_mismatch")}
    assert all(i["venue"] == "A" for i in issues)
    assert next(i for i in issues if i["reason"] == "range_mismatch")["blocking"]


def test_target_aliases_apply_only_to_the_selected_workbook(tmp_path, monkeypatch):
    import openpyxl
    september = tmp_path / "2026-09_目标.xlsx"
    august = tmp_path / "2026-08_目标.xlsx"
    september.with_suffix(".aliases.json").write_text(
        json.dumps({"福州A广场": "福州A店"}), encoding="utf-8"
    )
    monkeypatch.setattr(targets, "find_target_file", lambda month: str(september if month == "2026-09" else august))
    workbook = mock.Mock()
    workbook.active.iter_rows.side_effect = lambda **kwargs: iter([
        ("场地", "最终核定目标"), ("福州A广场", 200000), ("合计", 200000)
    ])
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *args, **kwargs: workbook)
    assert targets.load_store_targets("2026-09") == {"福州A店": 200000}
    assert targets.load_store_targets("2026-08") == {"福州A广场": 200000}
    workbook.active.iter_rows.side_effect = lambda **kwargs: iter([
        ("场地", "最终核定目标"), ("福州A广场", 200000), ("福州A店", 100000)
    ])
    with pytest.raises(ValueError, match="重复门店"):
        targets.load_store_targets("2026-09")


def test_kowloon_september_never_reads_august_sheet():
    from crawlers import coin_exchange_crawler as coin
    with mock.patch.object(coin, "_fetch_kdocs_amounts", side_effect=AssertionError("must not read August sheet")):
        assert coin.main("2026-09-07") == []
        assert coin.main("2026-09-07", [3770]) == []
    assert coin.main("2026-08-07", [3770]) == [{"场地": coin.DEFAULT_VENUE, "兑币机收款": 3770.0}]
